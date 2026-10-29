import argparse
import streamlit as st
import html
import io
import json
import logging
import math
import numpy as np
import openpyxl
import openpyxl.styles
import os
import re
import requests
import sys
import subprocess
import threading
import time
import truststore
from datetime import datetime, time as datetime_time
from dotenv import load_dotenv
from telegram_inbox import list_conversations, list_messages, record_message
from stock_data import fetch_price_board
from stock_history import (
    calculate_intraday_signals,
    get_average_volumes,
    get_history_summary,
    get_latest_daily_quotes,
    get_latest_signal_snapshot,
    get_recent_signals,
    has_final_daily_close,
    import_csv_files,
    record_daily_close,
)

# Dung kho chung chi cua Windows de xac thuc SSL (can khi mang cong ty co proxy kiem tra SSL)
truststore.inject_into_ssl()

# Cho phep chay bang "python app.py" (nut ▶ cua VS Code): tu khoi dong lai qua "streamlit run"
if __name__ == "__main__":
    from streamlit import runtime
    if not runtime.exists():
        from streamlit.web import cli as stcli
        parser = argparse.ArgumentParser()
        parser.add_argument("--port", "-port", type=int, default=8501)
        args = parser.parse_args()
        sys.argv = ["streamlit", "run", "--server.port", str(args.port), __file__]
        sys.exit(stcli.main())

POLL_SECONDS = 3  # Chu ky lay du lieu moi va lam moi bang
PRICE_UPDATE_TIMES = (
    (9, 17),
    (9, 32),
    (10, 2),
    (10, 32),
    (11, 2),
    (11, 32),
    (13, 2),
    (13, 32),
    (14, 2),
    (14, 17),
    (14, 32),
    (14, 47),
    (15, 2),
)
PRICE_UPDATE_SCHEDULE_TEXT = ", ".join(
    f"{hour:02d}:{minute:02d}" for hour, minute in PRICE_UPDATE_TIMES
)
BOARD_RESUME_TIME = datetime_time(8, 45)
BOARD_FREEZE_TIME = datetime_time(15, 30)
STOCK_DETAIL_URL = "https://24hmoney.vn/stock/{ticker}"  # Trang mo khi bam vao 1 ma


def is_board_frozen(now):
    if now.weekday() >= 5:
        return True
    return now.time() < BOARD_RESUME_TIME or now.time() >= BOARD_FREEZE_TIME


def is_after_close(now):
    return now.weekday() < 5 and now.time() >= BOARD_FREEZE_TIME

# 1. Cau hinh trang Streamlit
st.set_page_config(
    page_title="Bang Gia Chung Khoan Phan Nganh - FireAnt Style",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# 2. Dinh nghia CSS custom cho giao dien Dark Mode chuan chung khoan
st.markdown("""
    <style>
    .stApp {
        background-color: #0d1117;
        color: #f0f6fc;
    }
    .stApp [data-testid="stCaptionContainer"] {
        color: #8b949e;
    }
    /* Bo cuc gon de toan bo cac nganh vua 1 man hinh 1920x1080 */
    header[data-testid="stHeader"] {
        display: none;
    }
    .block-container {
        padding: 0.5rem 1rem 0.5rem 1rem;
        max-width: 100%;
    }
    .block-container [data-testid="stVerticalBlock"] {
        gap: 0.25rem;
    }
    .app-title {
        font-size: 22px;
        font-weight: 700;
        color: #f0f6fc;
        line-height: 1.2;
    }
    .board {
        column-count: 6;      /* 6-8 cot (dat trong render_board), trinh duyet tu can bang chieu cao cac cot */
        column-gap: 8px;
    }
    .sector-card {
        background-color: #161b22;
        border: 1px solid #30363d;
        border-radius: 6px;
        padding: 4px 5px;
        margin-bottom: 6px;
        break-inside: avoid;   /* khong cat 1 nganh ra 2 cot */
    }
    .sector-header {
        font-weight: 700;
        color: #f0f6fc;
        font-size: 13px;
        letter-spacing: 0.5px;
        border-bottom: 1px solid #30363d;
        padding: 1px 2px 3px 2px;
        margin-bottom: 3px;
        display: flex;
        justify-content: space-between;
        align-items: center;
    }
    .stock-row {
        font-family: 'Roboto Mono', monospace, monospace;
        font-size: 13px;
        font-weight: 600;
        line-height: 19px;
        display: flex;
        justify-content: space-between;
        align-items: center;
        padding: 1px 6px;
        border-radius: 3px;
        margin-bottom: 1px;
        text-decoration: none !important;  /* moi dong la 1 link sang 24hmoney - bo gach chan */
    }
    .stock-row:hover {
        filter: brightness(1.5);
        outline: 1px solid rgba(240, 246, 252, 0.4);
    }
    /* !important de mau theo gia khong bi mau link mac dinh cua Streamlit ghi de */
    .ceil { background-color: rgba(156, 39, 176, 0.25); color: #e040fb !important; border-left: 3px solid #e040fb; }
    .up { background-color: rgba(46, 125, 50, 0.25); color: #4caf50 !important; border-left: 3px solid #4caf50; }
    .ref { background-color: rgba(245, 124, 0, 0.2); color: #ffb74d !important; border-left: 3px solid #ffb74d; }
    .down { background-color: rgba(198, 40, 40, 0.25); color: #ef5350 !important; border-left: 3px solid #ef5350; }
    .floor { background-color: rgba(2, 136, 209, 0.25); color: #29b6f6 !important; border-left: 3px solid #29b6f6; }
    </style>
""", unsafe_allow_html=True)

# 3. Danh sach cac nhom nganh va ma co phieu tuong ung
#    Danh sach mac dinh - dung khi chua co file sectors.json (nguoi dung sua qua nut "Quan ly nganh")
DEFAULT_SECTORS = {
    "NGAN HANG": ["EIB", "STB", "TPB", "MSB", "VIB", "BID", "CTG", "MBB", "TCB", "VCB", "ACB", "VPB", "HDB", "NVB",  # NVB = ma CK cua NCB
                  "ABB", "BVB", "LPB"],
    "CHUNG KHOAN": ["ORS", "VIX", "MBS", "VDS", "SHS", "VCI", "BSI", "SSI", "VND", "CTS", "FTS", "HCM",
                    "TCX", "VCK", "VPX", "AAS", "APS"],
    "BAT DONG SAN": ["HDC", "CEO", "NLG", "DIG", "VHM", "HDG", "VIC", "DXG", "PDR", "NVL", "KDH", "TCH", "HHS"],
    "VLXD - DAU TU CONG": ["C4G", "G36", "HBC", "LCG", "CII", "HHV", "FCN", "HUT", "CTD", "KSB", "VCG", "HT1"],
    "BDS KCN - CAO SU": ["SZC", "DRC", "TDC", "DPR", "KBC", "PHR", "LHG", "VGC", "TIP", "GVR", "IDC", "BCM"],
    "DIEN NUOC": ["GEE", "GEX", "GEL", "PC1", "QTP", "GEG", "REE", "NT2", "POW", "BWE", "PPC", "TDM"],
    "THUC PHAM": ["DBC", "PAN", "HNG", "QNS", "HAG", "BAF", "SAB", "MSN", "VNM", "KDC", "SBT", "SLS", "NAF"],
    "XUAT KHAU": ["ASM", "TCM", "VGT", "GIL", "TNG", "VHC", "ANV", "IDI", "STK", "FMC"],
    "DAU KHI": ["OIL", "PVC", "PLX", "BSR", "GAS", "PVD", "PVB", "PVS", "PVT"],
    "CANG BIEN": ["VSC", "SGP", "HAH", "GMD", "DVP", "PHP", "PVP", "VOS", "VIP", "DXP"],
    "HOA CHAT - PHAN BON": ["DPM", "DGC", "DDV", "CSV", "DCM", "LAS", "BFC"],
    "THEP": ["NKG", "HSG", "HPG", "TLH", "SMC", "TVN"],
    "CNTT": ["FPT", "FOX", "FOC", "CMG", "VGI", "ELC", "CTR", "MFS", "ICT"],
    "KHOANG SAN": ["MSR", "BKC", "KSV"],
    "BAN LE": ["FRT", "PET", "MWG", "DMX", "PNJ", "DGW"]
}

# File luu danh sach nganh do nguoi dung chinh sua: canh StockBoard.exe khi chay ban .exe, canh app.py khi chay code
DATA_DIR = os.path.dirname(sys.executable) if getattr(sys, "frozen", False) else os.path.dirname(os.path.abspath(__file__))
SECTORS_FILE = os.path.join(DATA_DIR, "sectors.json")
TELEGRAM_ALERTS_FILE = os.path.join(DATA_DIR, "telegram_alerts.json")
TELEGRAM_BOT_STATUS_FILE = os.path.join(DATA_DIR, ".telegram_bot_status")
TELEGRAM_BOT_STOP_FILE = os.path.join(DATA_DIR, ".telegram_bot_stop")
TELEGRAM_BOT_SCRIPT = os.path.join(DATA_DIR, "telegram_bot.py")
TELEGRAM_BOT_LOG = os.path.join(DATA_DIR, "telegram_bot.log")
TELEGRAM_INBOX_DATABASE = os.path.join(DATA_DIR, "telegram_inbox.db")
STOCK_HISTORY_DATABASE = os.path.join(DATA_DIR, "data", "stock_history.sqlite3")
load_dotenv(os.path.join(DATA_DIR, "Phong_telegram.env"))
TELEGRAM_BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")


def load_sectors():
    try:
        with open(SECTORS_FILE, encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict) and all(isinstance(v, list) for v in data.values()):
            return {str(k): [str(t) for t in v] for k, v in data.items()}
    except (OSError, ValueError):
        pass  # Chua co file hoac file hong -> dung danh sach mac dinh
    return {k: list(v) for k, v in DEFAULT_SECTORS.items()}


def save_sectors(sectors):
    # Ghi ra file tam roi doi ten, tranh hong file neu bi ngat giua chung
    tmp_path = SECTORS_FILE + ".tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump(sectors, f, ensure_ascii=False, indent=2)
    os.replace(tmp_path, SECTORS_FILE)


def parse_telegram_chat_ids(value):
    values = value if isinstance(value, (list, tuple)) else str(value or "").split(",")
    return list(dict.fromkeys(str(item).strip() for item in values if str(item).strip()))


def load_telegram_alert_settings():
    settings = {
        "enabled": True,
        "threshold": 2.0,
        "duration_minutes": 15,
        "signal_interval_minutes": 30,
        "volume_lookback_sessions": 10,
        "price_update_enabled": False,
        "buy_sell_enabled": True,
        "ma_short_period": 20,
        "ma_long_period": 50,
        "rsi_level1": 50.0,
        "rsi_level2": 70.0,
        "chat_ids": parse_telegram_chat_ids(
            os.getenv("TELEGRAM_CHAT_IDS") or os.getenv("TELEGRAM_CHAT_ID", "")
        ),
    }
    try:
        with open(TELEGRAM_ALERTS_FILE, encoding="utf-8") as f:
            saved = json.load(f)
        if isinstance(saved, dict):
            if isinstance(saved.get("enabled"), bool):
                settings["enabled"] = saved["enabled"]
            if isinstance(saved.get("price_update_enabled"), bool):
                settings["price_update_enabled"] = saved["price_update_enabled"]
            if isinstance(saved.get("buy_sell_enabled"), bool):
                settings["buy_sell_enabled"] = saved["buy_sell_enabled"]
            elif isinstance(saved.get("enabled"), bool):
                settings["buy_sell_enabled"] = saved["enabled"]
            threshold = float(saved.get("threshold", settings["threshold"]))
            if math.isfinite(threshold) and 0 < threshold <= 100:
                settings["threshold"] = threshold
            duration = int(saved.get("duration_minutes", settings["duration_minutes"]))
            if 1 <= duration <= 1440:
                settings["duration_minutes"] = duration
            signal_interval = int(
                saved.get("signal_interval_minutes", settings["signal_interval_minutes"])
            )
            if 5 <= signal_interval <= 360:
                settings["signal_interval_minutes"] = signal_interval
            volume_lookback = int(
                saved.get("volume_lookback_sessions", settings["volume_lookback_sessions"])
            )
            if 10 <= volume_lookback <= 100 and volume_lookback % 5 == 0:
                settings["volume_lookback_sessions"] = volume_lookback
            ma_short_period = int(saved.get("ma_short_period", settings["ma_short_period"]))
            ma_long_period = int(saved.get("ma_long_period", settings["ma_long_period"]))
            if 2 <= ma_short_period < ma_long_period <= 500:
                settings["ma_short_period"] = ma_short_period
                settings["ma_long_period"] = ma_long_period
            rsi_level1 = float(saved.get("rsi_level1", settings["rsi_level1"]))
            rsi_level2 = float(saved.get("rsi_level2", settings["rsi_level2"]))
            if math.isfinite(rsi_level1) and math.isfinite(rsi_level2):
                if 0 <= rsi_level1 < rsi_level2 <= 100:
                    settings["rsi_level1"] = rsi_level1
                    settings["rsi_level2"] = rsi_level2
            if "chat_ids" in saved:
                settings["chat_ids"] = parse_telegram_chat_ids(saved["chat_ids"])
    except (OSError, TypeError, ValueError):
        pass
    return settings


def save_telegram_alert_settings(
    enabled,
    threshold,
    duration_minutes,
    signal_interval_minutes,
    volume_lookback_sessions,
    price_update_enabled,
    buy_sell_enabled,
    ma_short_period,
    ma_long_period,
    rsi_level1,
    rsi_level2,
    chat_ids,
):
    tmp_path = TELEGRAM_ALERTS_FILE + ".tmp"
    with open(tmp_path, "w", encoding="utf-8") as f:
        json.dump({
            "enabled": enabled,
            "threshold": threshold,
            "duration_minutes": duration_minutes,
            "signal_interval_minutes": signal_interval_minutes,
            "volume_lookback_sessions": volume_lookback_sessions,
            "price_update_enabled": price_update_enabled,
            "buy_sell_enabled": buy_sell_enabled,
            "ma_short_period": ma_short_period,
            "ma_long_period": ma_long_period,
            "rsi_level1": rsi_level1,
            "rsi_level2": rsi_level2,
            "chat_ids": chat_ids,
        }, f, indent=2)
    os.replace(tmp_path, TELEGRAM_ALERTS_FILE)


def get_telegram_bot_status():
    try:
        heartbeat_age = time.time() - os.path.getmtime(TELEGRAM_BOT_STATUS_FILE)
    except OSError:
        return "Stopped"
    return "Running" if 0 <= heartbeat_age < 20 else "Stopped"


def start_telegram_bot():
    if get_telegram_bot_status() == "Running":
        return False, "Telegram bot is already running."
    if not TELEGRAM_BOT_TOKEN:
        return False, "TELEGRAM_BOT_TOKEN is missing from Phong_telegram.env."
    if getattr(sys, "frozen", False):
        return False, "Start the bot with Python when using the packaged StockBoard executable."
    if not os.path.isfile(TELEGRAM_BOT_SCRIPT):
        return False, f"Could not find {TELEGRAM_BOT_SCRIPT}."

    try:
        try:
            os.remove(TELEGRAM_BOT_STOP_FILE)
        except FileNotFoundError:
            pass
        with open(TELEGRAM_BOT_LOG, "a", encoding="utf-8") as log_file:
            subprocess.Popen(
                [sys.executable, TELEGRAM_BOT_SCRIPT],
                cwd=DATA_DIR,
                stdin=subprocess.DEVNULL,
                stdout=log_file,
                stderr=subprocess.STDOUT,
                close_fds=True,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
            )
    except OSError as error:
        logging.exception("Could not start Telegram bot")
        return False, f"Could not start Telegram bot: {error}"
    return True, "Telegram bot start requested. Its status will update shortly."


def stop_telegram_bot():
    if get_telegram_bot_status() != "Running":
        return False, "Telegram bot is not running."
    try:
        with open(TELEGRAM_BOT_STOP_FILE, "w", encoding="ascii") as stop_request:
            stop_request.write("stop")
    except OSError as error:
        logging.exception("Could not request Telegram bot shutdown")
        return False, f"Could not stop Telegram bot: {error}"

    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        if get_telegram_bot_status() != "Running":
            return True, "Telegram bot stopped."
        time.sleep(0.2)

    try:
        os.remove(TELEGRAM_BOT_STOP_FILE)
    except FileNotFoundError:
        pass
    except OSError:
        logging.exception("Could not clear Telegram bot stop request")
    return False, "Bot did not stop; restart it to enable web start/stop control."


SECTOR_MAP = load_sectors()
ALL_TICKERS = list(dict.fromkeys(t for tickers in SECTOR_MAP.values() for t in tickers))

SESSION_LABELS = {
    "ATO": "Phien ATO", "LO": "Phien lien tuc", "LO_MORNING": "Phien sang", "LO_AFTERNOON": "Phien chieu",
    "BREAK": "Nghi trua", "INTERMISSION": "Nghi trua",
    "ATC": "Phien ATC", "PT": "Thoa thuan", "ENDED": "Da dong cua",
}


# 4. Kho du lieu realtime dung chung + Background Thread cap nhat lien tuc
#    (Thread ngam KHONG duoc dung st.session_state - chi ghi vao kho nay, co khoa de an toan luong)
class LiveQuoteStore:
    def __init__(self):
        self.lock = threading.Lock()
        self.wake = threading.Event()
        self.ready = threading.Event()  # duoc bat sau lan lay du lieu dau tien
        self.tickers = list(ALL_TICKERS)  # danh sach ma can lay, cap nhat khi sua SECTOR_MAP
        self.quotes = get_latest_daily_quotes(STOCK_HISTORY_DATABASE)
        # Khi app mo ngoai gio giao dich, hien thi gia dong cua SQLite gan nhat.
        self.last_update = None   # thoi diem lay du lieu thanh cong gan nhat
        self.session = None       # trang thai phien giao dich
        self.error = None         # loi gan nhat (None neu dang ket noi tot)
        self.threshold_alert_date = None
        self.threshold_alert_slots = set()
        alert_settings = load_telegram_alert_settings()
        self.telegram_enabled = alert_settings["enabled"]
        self.telegram_threshold = alert_settings["threshold"]
        self.telegram_duration_minutes = alert_settings["duration_minutes"]
        self.telegram_signal_interval_minutes = alert_settings[
            "signal_interval_minutes"
        ]
        self.volume_lookback_sessions = alert_settings["volume_lookback_sessions"]
        self.price_update_enabled = alert_settings["price_update_enabled"]
        self.buy_sell_enabled = alert_settings["buy_sell_enabled"]
        self.ma_short_period = alert_settings["ma_short_period"]
        self.ma_long_period = alert_settings["ma_long_period"]
        self.rsi_level1 = alert_settings["rsi_level1"]
        self.rsi_level2 = alert_settings["rsi_level2"]
        self.telegram_chat_ids = alert_settings["chat_ids"]
        self.signal_alert_date = None
        self.signal_alerted_slots = set()
        self.price_update_date = None
        self.price_update_slots = set()
        self.price_update_last_prices = None

    def snapshot(self):
        with self.lock:
            return dict(self.quotes), self.last_update, self.session, self.error

    def refresh_now(self):
        self.wake.set()

    def signal_settings(self):
        with self.lock:
            return (
                self.buy_sell_enabled,
                self.ma_short_period,
                self.ma_long_period,
                self.rsi_level1,
                self.rsi_level2,
            )

    def update_signal_settings(
        self, enabled, ma_short_period, ma_long_period, rsi_level1, rsi_level2
    ):
        with self.lock:
            self.buy_sell_enabled = enabled
            self.ma_short_period = ma_short_period
            self.ma_long_period = ma_long_period
            self.rsi_level1 = rsi_level1
            self.rsi_level2 = rsi_level2

    def price_update_settings(self):
        with self.lock:
            return (
                self.price_update_enabled,
                list(self.telegram_chat_ids),
            )

    def volume_settings(self):
        with self.lock:
            return self.volume_lookback_sessions

    def update_volume_settings(self, lookback_sessions):
        with self.lock:
            self.volume_lookback_sessions = lookback_sessions

    def update_price_update_settings(self, enabled):
        with self.lock:
            self.price_update_enabled = enabled

    def alert_settings(self):
        with self.lock:
            return (
                self.telegram_enabled,
                self.telegram_threshold,
                self.telegram_duration_minutes,
                self.telegram_signal_interval_minutes,
                list(self.telegram_chat_ids),
            )

    def update_alert_settings(
        self, enabled, threshold, duration_minutes, signal_interval_minutes, chat_ids
    ):
        with self.lock:
            self.telegram_enabled = enabled
            self.telegram_threshold = threshold
            self.telegram_duration_minutes = duration_minutes
            self.telegram_signal_interval_minutes = signal_interval_minutes
            self.telegram_chat_ids = list(chat_ids)


def send_telegram_message(chat_id, message):
    if not TELEGRAM_BOT_TOKEN:
        return False, "The bot token is not configured."
    try:
        response = requests.post(
            f"https://api.telegram.org/bot{TELEGRAM_BOT_TOKEN}/sendMessage",
            data={"chat_id": chat_id, "text": message},
            timeout=5,
        )
    except requests.RequestException:
        logging.warning("Could not connect to Telegram Bot API.")
        return False, "Could not reach Telegram; check the network connection."

    try:
        result = response.json()
    except ValueError:
        result = {}
    if response.ok and isinstance(result, dict) and result.get("ok"):
        record_message(TELEGRAM_INBOX_DATABASE, chat_id, "outgoing", message)
        return True, None

    reason = result.get("description", f"Telegram returned HTTP {response.status_code}.")
    logging.warning("Telegram rejected a message: %s", reason)
    return False, reason


def volume_relative_to_average(ticker, quotes, average_volumes):
    quote = quotes[ticker]
    volume = quote.get("volume")
    if volume is None:
        volume = quote.get("vol", 0) * 1000
    average_volume = average_volumes.get(ticker)
    if average_volume is None or average_volume <= 0:
        return volume, "N/A"
    return volume, f"{volume / average_volume * 100:.0f}%"


def format_ticker_alert_sections(
    strong_tickers,
    weak_tickers,
    quotes,
    threshold,
    average_volumes,
    lookback_sessions,
):
    def format_row(ticker):
        quote = quotes[ticker]
        volume, relative_volume = volume_relative_to_average(
            ticker, quotes, average_volumes
        )
        return (
            f"{ticker:<7} {quote['price']:>8.2f} {quote['pct']:>+7.2f}% "
            f"{volume:>12,.0f} {relative_volume:>8}"
        )

    strong_rows = "\n".join(format_row(ticker) for ticker in strong_tickers) or "None"
    weak_rows = "\n".join(format_row(ticker) for ticker in weak_tickers) or "None"
    header = (
        f"{'Ma CK':<7} {'Gia':>8} {'Thay doi':>8} {'KLGD':>12} "
        f"{'Vol/TB ' + str(lookback_sessions) + 'p':>8}"
    )
    return (
        f"1. Co Phieu Manh (>={threshold:g}%) :\n"
        "----------------------------\n"
        f"{header}\n"
        f"{strong_rows}\n\n"
        f"2. Co Phieu Yeu (<=-{threshold:g}%) :\n"
        "----------------------------\n"
        f"{header}\n"
        f"{weak_rows}"
    )


def send_threshold_alerts(store, quotes, now):
    enabled, threshold, _, _, chat_ids = store.alert_settings()
    if not enabled or not TELEGRAM_BOT_TOKEN or not chat_ids:
        return

    slot = scheduled_price_update_slot(now)
    if slot is None:
        return

    volume_lookback_sessions = store.volume_settings()
    with store.lock:
        if store.threshold_alert_date != now.date():
            store.threshold_alert_date = now.date()
            store.threshold_alert_slots.clear()
        if slot in store.threshold_alert_slots:
            return
        store.threshold_alert_slots.add(slot)
        tickers = list(store.tickers)

    strong_tickers = [
        ticker for ticker in tickers
        if ticker in quotes and quotes[ticker].get("pct") is not None
        and quotes[ticker]["pct"] >= threshold
    ]
    weak_tickers = [
        ticker for ticker in tickers
        if ticker in quotes and quotes[ticker].get("pct") is not None
        and quotes[ticker]["pct"] <= -threshold
    ]
    if not strong_tickers and not weak_tickers:
        return

    average_volumes = get_average_volumes(
        STOCK_HISTORY_DATABASE,
        tickers,
        now.date(),
        volume_lookback_sessions,
    )
    message = (
        f"StockBoard daily change snapshot ({now:%Y-%m-%d %H:%M}; "
        f"threshold: +/-{threshold:g}%)\n\n"
        + format_ticker_alert_sections(
            strong_tickers,
            weak_tickers,
            quotes,
            threshold,
            average_volumes,
            volume_lookback_sessions,
        )
    )
    for chat_id in chat_ids:
        sent, error = send_telegram_message(chat_id, message)
        if not sent:
            logging.warning("Telegram alert to a recipient failed: %s", error)


def scheduled_signal_slot(now, interval_minutes):
    if now.weekday() >= 5:
        return None
    minute_of_day = now.hour * 60 + now.minute
    if (
        9 * 60 <= minute_of_day <= 15 * 60
        and (minute_of_day - 9 * 60) % interval_minutes == 0
    ):
        return minute_of_day
    if minute_of_day == 15 * 60 + 30:
        return minute_of_day
    return None


def scheduled_price_update_slot(now):
    if now.weekday() >= 5 or (now.hour, now.minute) not in PRICE_UPDATE_TIMES:
        return None
    return now.hour * 60 + now.minute


def send_scheduled_signal_alerts(store, quotes, now):
    signal_settings = store.signal_settings()
    signal_enabled, ma_short_period, ma_long_period, rsi_level1, rsi_level2 = signal_settings
    _, _, _, interval_minutes, chat_ids = store.alert_settings()
    if not signal_enabled or not TELEGRAM_BOT_TOKEN or not chat_ids:
        return
    slot = scheduled_signal_slot(now, interval_minutes)
    if slot is None:
        return

    with store.lock:
        if store.signal_alert_date != now.date():
            store.signal_alert_date = now.date()
            store.signal_alerted_slots.clear()
        if slot in store.signal_alerted_slots:
            return
        store.signal_alerted_slots.add(slot)
        tickers = list(store.tickers)

    try:
        signals = calculate_intraday_signals(
            STOCK_HISTORY_DATABASE,
            quotes,
            tickers,
            now.date(),
            ma_short_period,
            ma_long_period,
            rsi_level1,
            rsi_level2,
        )
    except Exception:
        logging.exception("Could not calculate scheduled intraday signals")
        return
    if not signals:
        return

    visible_signals = signals[:50]
    message = (
        f"StockBoard BUY/SELL signals ({now:%Y-%m-%d %H:%M})\n"
        f"{len(signals)} signal(s); showing {len(visible_signals)}:\n\n"
        + "\n".join(
            f"{item['ticker']} {item['signal']} | Price {item['price']:.2f} | "
            f"MA{item['ma_short_period']} {item['ma20']:.2f} | "
            f"MA{item['ma_long_period']} {item['ma50']:.2f} | "
            f"RSI14 {item['rsi14']:.1f}"
            for item in visible_signals
        )
    )
    for chat_id in chat_ids:
        sent, error = send_telegram_message(chat_id, message)
        if not sent:
            logging.warning(
                "Scheduled signal alert to %s failed: %s", chat_id, error
            )


def send_periodic_price_updates(store, quotes, now):
    enabled, chat_ids = store.price_update_settings()
    if not enabled or not TELEGRAM_BOT_TOKEN or not chat_ids:
        return
    slot = scheduled_price_update_slot(now)
    volume_lookback_sessions = store.volume_settings()

    with store.lock:
        if store.price_update_date != now.date():
            store.price_update_date = now.date()
            store.price_update_slots.clear()
            store.price_update_last_prices = None
        tickers = list(store.tickers)
        current_prices = {}
        for ticker in tickers:
            quote = quotes.get(ticker)
            if not quote:
                continue
            try:
                price = float(quote["price"])
            except (KeyError, TypeError, ValueError):
                continue
            if math.isfinite(price) and price > 0:
                current_prices[ticker] = price

        if store.price_update_last_prices is None and current_prices:
            store.price_update_last_prices = current_prices
        if slot is None or slot in store.price_update_slots:
            return
        store.price_update_slots.add(slot)
        previous_prices = store.price_update_last_prices
        if previous_prices is None:
            return
        changed = [
            (ticker, previous_prices[ticker], price)
            for ticker, price in current_prices.items()
            if ticker in previous_prices
            and not math.isclose(price, previous_prices[ticker], rel_tol=0, abs_tol=1e-9)
        ]
        if not changed:
            store.price_update_last_prices = current_prices
            return

    average_volumes = get_average_volumes(
        STOCK_HISTORY_DATABASE,
        tickers,
        now.date(),
        volume_lookback_sessions,
    )
    lines = [
        f"{ticker}: {old_price:.2f} -> {new_price:.2f} "
        f"({(new_price - old_price) / old_price * 100:+.2f}%) | "
        f"Vol {volume_relative_to_average(ticker, quotes, average_volumes)[0]:,.0f} | "
        f"Vol/TB {volume_lookback_sessions}p "
        f"{volume_relative_to_average(ticker, quotes, average_volumes)[1]}"
        for ticker, old_price, new_price in changed[:80]
    ]
    message = (
        f"StockBoard price changes ({now:%Y-%m-%d %H:%M})\n"
        f"{len(changed)} ticker(s) changed since the previous update; "
        f"showing {len(lines)}:\n\n"
        + "\n".join(lines)
    )
    all_sent = True
    for chat_id in chat_ids:
        sent, error = send_telegram_message(chat_id, message)
        if not sent:
            all_sent = False
            logging.warning("Periodic price update to %s failed: %s", chat_id, error)
    if all_sent:
        with store.lock:
            store.price_update_last_prices = current_prices


def send_telegram_test(
    chat_ids, threshold, price_update_enabled,
    signal_interval_minutes, volume_lookback_sessions,
    buy_sell_settings, tickers, quotes, updated_at,
):
    buy_sell_enabled, ma_short_period, ma_long_period, rsi_level1, rsi_level2 = (
        buy_sell_settings
    )
    message = (
        "StockBoard test message. Telegram delivery is working.\n"
        f"Threshold snapshot: +/-{threshold:g}% at {PRICE_UPDATE_SCHEDULE_TEXT}.\n"
        f"BUY/SELL signal interval: {signal_interval_minutes} min; "
        f"volume average lookback {volume_lookback_sessions} sessions.\n"
        "Price change schedule runs on weekdays using app local time.\n"
    )
    if not buy_sell_enabled:
        message += "BUY/SELL Telegram sending is disabled."
    else:
        message += (
            f"BUY/SELL rule: MA{ma_short_period}/{ma_long_period}; "
            f"RSI levels {rsi_level1:g}/{rsi_level2:g}."
        )
    if not quotes:
        message += "\n\nLive ticker snapshot is not available."
    else:
        strong_tickers = sorted(
            (ticker for ticker in tickers if ticker in quotes and quotes[ticker]["pct"] > threshold),
            key=lambda ticker: abs(quotes[ticker]["pct"]),
            reverse=True,
        )
        weak_tickers = sorted(
            (ticker for ticker in tickers if ticker in quotes and quotes[ticker]["pct"] < -threshold),
            key=lambda ticker: abs(quotes[ticker]["pct"]),
            reverse=True,
        )
        snapshot_time = updated_at.strftime("%Y-%m-%d %H:%M:%S") if updated_at else "unknown"
        snapshot_date = updated_at.date() if updated_at else datetime.now().date()
        average_volumes = get_average_volumes(
            STOCK_HISTORY_DATABASE,
            tickers,
            snapshot_date,
            volume_lookback_sessions,
        )
        message += f"\n\nLive ticker snapshot (updated {snapshot_time}):\n"
        message += format_ticker_alert_sections(
            strong_tickers,
            weak_tickers,
            quotes,
            threshold,
            average_volumes,
            volume_lookback_sessions,
        )

    if buy_sell_enabled:
        signal_date, signals = get_latest_signal_snapshot(
            STOCK_HISTORY_DATABASE,
            tickers=tickers,
        )
        message += "\n\nLatest daily BUY/SELL signals"
        if signal_date is None:
            message += ": no signals have been recorded yet."
        elif not signals:
            message += f" ({signal_date}): no BUY/SELL signals."
        else:
            message += f" ({signal_date}; up to {len(signals)} shown):\n"
            message += "\n".join(
                f"{item['ticker']} {item['signal']} | "
                f"MA{item['ma_short_period']} {item['ma20']:.2f} | "
                f"MA{item['ma_long_period']} {item['ma50']:.2f} | "
                f"RSI14 {item['rsi14']:.1f}"
                for item in signals
            )

    successful, failed = [], []
    for chat_id in chat_ids:
        sent, error = send_telegram_message(chat_id, message)
        if sent:
            successful.append(chat_id)
        else:
            failed.append((chat_id, error))
    return successful, failed


@st.fragment(run_every=5)
def render_telegram_inbox():
    conversations = list_conversations(TELEGRAM_INBOX_DATABASE)
    if not conversations:
        st.caption("No Telegram conversations yet.")
        return

    conversation_by_id = {str(item["chat_id"]): item for item in conversations}
    chat_ids = list(conversation_by_id)
    if st.session_state.get("telegram_inbox_chat") not in chat_ids:
        st.session_state["telegram_inbox_chat"] = chat_ids[0]

    selected_chat_id = st.selectbox(
        "Conversation",
        options=chat_ids,
        format_func=lambda chat_id: (
            f"{conversation_by_id[chat_id]['display_name'] or chat_id}"
            + (
                f" (@{conversation_by_id[chat_id]['username']})"
                if conversation_by_id[chat_id]["username"]
                else ""
            )
        ),
        key="telegram_inbox_chat",
    )
    with st.container(height=320, border=True):
        for item in list_messages(TELEGRAM_INBOX_DATABASE, selected_chat_id):
            role = "user" if item["direction"] == "incoming" else "assistant"
            with st.chat_message(role):
                created_at = datetime.fromisoformat(item["created_at"]).astimezone()
                st.caption(created_at.strftime("%Y-%m-%d %H:%M:%S"))
                st.write(item["text"])

    with st.form(f"telegram_inbox_reply_{selected_chat_id}", clear_on_submit=True):
        reply_text = st.text_area("Reply", key=f"telegram_reply_text_{selected_chat_id}")
        send_reply = st.form_submit_button("Send reply", width="stretch")
    if send_reply:
        reply_text = reply_text.strip()
        if not reply_text:
            st.error("Enter a reply first.")
        else:
            sent, error = send_telegram_message(selected_chat_id, reply_text)
            if sent:
                st.success("Reply sent.")
            else:
                st.error(f"Reply failed: {error}")


def poll_forever(store):
    while True:
        started = time.monotonic()
        try:
            now = datetime.now()
            if is_board_frozen(now):
                # After the close, keep polling only until the final daily snapshot
                # reaches SQLite. Before open and on weekends, do not call the API.
                if is_after_close(now) and not has_final_daily_close(
                    STOCK_HISTORY_DATABASE, now.date()
                ):
                    quotes, session = fetch_price_board(store.tickers)
                    now = datetime.now()
                    _, ma_short_period, ma_long_period, rsi_level1, rsi_level2 = (
                        store.signal_settings()
                    )
                    captured, _ = record_daily_close(
                        STOCK_HISTORY_DATABASE,
                        quotes,
                        now,
                        session,
                        ma_short_period,
                        ma_long_period,
                        rsi_level1,
                        rsi_level2,
                    )
                    if captured or has_final_daily_close(
                        STOCK_HISTORY_DATABASE, now.date()
                    ):
                        with store.lock:
                            store.quotes = quotes
                            store.session = session
                            store.last_update = now
                            store.error = None
                    elif not store.quotes:
                        with store.lock:
                            store.quotes = quotes
                            store.session = session
                            store.last_update = now
                            store.error = None
                # During the closed period the UI quote store is deliberately
                # unchanged; the next weekday API refresh starts at 08:45.
            else:
                quotes, session = fetch_price_board(store.tickers)
                now = datetime.now()
                _, ma_short_period, ma_long_period, rsi_level1, rsi_level2 = (
                    store.signal_settings()
                )
                try:
                    record_daily_close(
                        STOCK_HISTORY_DATABASE,
                        quotes,
                        now,
                        session,
                        ma_short_period,
                        ma_long_period,
                        rsi_level1,
                        rsi_level2,
                    )
                except Exception:
                    logging.exception("Could not capture daily stock history")
                send_threshold_alerts(store, quotes, now)
                send_periodic_price_updates(store, quotes, now)
                send_scheduled_signal_alerts(store, quotes, now)
                with store.lock:
                    store.quotes = quotes
                    store.session = session
                    store.last_update = now
                    store.error = None
        except Exception as e:
            with store.lock:
                store.error = str(e)
        store.ready.set()
        # Tru thoi gian goi API de giu chu ky deu POLL_SECONDS giay
        store.wake.wait(max(0.0, POLL_SECONDS - (time.monotonic() - started)))
        store.wake.clear()


# Chi khoi chay 1 thread duy nhat, dung chung cho moi nguoi dung/tab dang mo
@st.cache_resource
def start_live_feed():
    store = LiveQuoteStore()
    threading.Thread(target=poll_forever, args=(store,), daemon=True).start()
    store.ready.wait(timeout=15)  # Cho lan lay du lieu dau tien de khong hien du lieu gia lap luc mo app
    return store


# 5. Du lieu gia lap - chi dung khi chua ket noi duoc nguon realtime
@st.cache_data(ttl=10)
def get_simulated_data():
    quotes = {}
    for ticker in ALL_TICKERS:
        pct_change = round(np.random.normal(0.8, 2.2), 2)
        if pct_change >= 6.8:
            css_cls = "ceil"
        elif pct_change > 0:
            css_cls = "up"
        elif pct_change == 0:
            css_cls = "ref"
        elif pct_change <= -6.8:
            css_cls = "floor"
        else:
            css_cls = "down"
        quotes[ticker] = {
            "price": round(np.random.uniform(10.0, 75.0), 2),
            "pct": pct_change,
            "vol": round(np.random.uniform(150, 3800)),
            "cls": css_cls,
        }
    return quotes


store = start_live_feed()
# Thread ngam chay lien tuc nen khong tu thay ma moi them vao SECTOR_MAP - cap nhat cho no o moi lan chay script
if store.tickers != ALL_TICKERS:
    store.tickers = list(ALL_TICKERS)
    store.refresh_now()

# 6. Quan ly nganh: them/xoa/doi ten nganh, them/xoa ma trong nganh (luu vao sectors.json)
def clean_tickers(items):
    """Chuan hoa ma: viet hoa, bo trung; cho phep dan nhieu ma 1 lan (vd: 'dxs, nha')"""
    result = []
    for item in items:
        for part in re.split(r"[\s,;]+", str(item)):
            ticker = part.strip().upper()
            if ticker and ticker not in result:
                result.append(ticker)
    return result


def find_unknown_tickers(tickers):
    """Tra ve cac ma khong co tren bang gia (go sai, da huy niem yet...)"""
    quotes, _ = fetch_price_board(tickers)
    return [t for t in tickers if t not in quotes]


def validate_and_save(sectors, new_tickers, success_msg):
    """Kiem tra ma moi voi bang gia roi luu. Tra ve thong bao loi, hoac None neu da luu thanh cong."""
    if new_tickers:
        try:
            unknown = find_unknown_tickers(new_tickers)
        except requests.RequestException as e:
            return f"Khong kiem tra duoc ma co phieu (loi ket noi): {e}"
        if unknown:
            return f"Khong tim thay ma tren bang gia: {', '.join(unknown)}. Kiem tra lai ma (vd: NCB phai nhap la NVB)."
    try:
        save_sectors(sectors)
    except OSError as e:
        return f"Khong luu duoc file {SECTORS_FILE}: {e}"
    st.session_state["mg_flash"] = success_msg
    st.rerun()


# Xuat/nhap danh sach nganh bang Excel: moi dong 1 ma, 2 cot | Sector | Ticker |
EXCEL_HEADERS = ("Sector", "Ticker")


def sectors_to_excel(sectors):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sectors"
    ws.append(EXCEL_HEADERS)
    for sector, tickers in sectors.items():
        for ticker in tickers:
            ws.append((sector, ticker))
    for cell in ws[1]:
        cell.font = openpyxl.styles.Font(bold=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 12
    ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def sectors_from_excel(file):
    """Doc file Excel 2 cot Sector | Ticker (sheet dau tien). O Sector de trong = cung nganh voi dong tren."""
    try:
        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
    except Exception as e:
        raise ValueError(f"Khong doc duoc file Excel: {e}")
    sectors, current = {}, None
    for i, row in enumerate(wb.worksheets[0].iter_rows(values_only=True)):
        sector = str(row[0]).strip().upper() if len(row) > 0 and row[0] is not None else ""
        cell = row[1] if len(row) > 1 else None
        if i == 0 and sector == "SECTOR":
            continue  # dong tieu de
        if sector:
            current = sector
        tickers = clean_tickers([cell]) if cell is not None else []
        if not tickers:
            continue
        if current is None:
            raise ValueError(f"Dong {i + 1}: co ma {tickers[0]} nhung chua co ten nganh (cot Sector).")
        bucket = sectors.setdefault(current, [])
        bucket.extend(t for t in tickers if t not in bucket)
    wb.close()
    if not sectors:
        raise ValueError("File khong co du lieu. Can 2 cot: Sector | Ticker.")
    return sectors


@st.dialog("⚙️ Quan ly nganh", width="medium")
def manage_sectors():
    sectors = load_sectors()
    tab_edit, tab_add, tab_excel = st.tabs(["Sua / xoa nganh", "Them nganh moi", "📊 Excel"])

    with tab_edit:
        if not sectors:
            st.info("Chua co nganh nao. Them nganh o the \"Them nganh moi\".")
        else:
            sel = st.selectbox("Nganh", list(sectors), key="mg_sector")
            new_name = st.text_input("Ten nganh", value=sel, key=f"mg_name_{sel}")
            tickers = st.multiselect(
                "Ma co phieu", options=sectors[sel], default=sectors[sel],
                accept_new_options=True, key=f"mg_tickers_{sel}",
                placeholder="Go ma moi roi nhan Enter",
            )
            st.caption("Bam **×** tren ma de xoa · Go ma moi (vd: DXS) roi nhan **Enter** de them · Bam **Luu** de ap dung")

            col_save, col_delete = st.columns(2)
            if col_save.button("💾 Luu thay doi", type="primary", width="stretch", key="mg_save"):
                name = new_name.strip().upper()
                cleaned = clean_tickers(tickers)
                if not name:
                    st.error("Ten nganh khong duoc de trong.")
                elif name != sel and name in sectors:
                    st.error(f"Da co nganh \"{name}\".")
                elif not cleaned:
                    st.error("Nganh can it nhat 1 ma. Muon bo ca nganh thi bam \"Xoa nganh\".")
                else:
                    # Giu nguyen vi tri nganh tren bang khi doi ten
                    updated = {(name if k == sel else k): (cleaned if k == sel else v) for k, v in sectors.items()}
                    added = [t for t in cleaned if t not in sectors[sel]]
                    error = validate_and_save(updated, added, f"Da luu nganh {name} ({len(cleaned)} ma)")
                    st.error(error)
            with col_delete.popover("🗑️ Xoa nganh", width="stretch"):
                st.write(f"Xoa nganh **{sel}** ({len(sectors[sel])} ma) khoi bang gia?")
                if st.button("Xac nhan xoa", type="primary", key="mg_delete"):
                    remaining = {k: v for k, v in sectors.items() if k != sel}
                    st.error(validate_and_save(remaining, [], f"Da xoa nganh {sel}"))

    with tab_add:
        name = st.text_input("Ten nganh moi", key="mg_new_name", placeholder="VD: HANG KHONG")
        raw = st.text_input("Ma co phieu (cach nhau boi dau phay)", key="mg_new_tickers", placeholder="VD: HVN, VJC, ACV")
        if st.button("➕ Them nganh", type="primary", key="mg_add"):
            name = name.strip().upper()
            cleaned = clean_tickers([raw])
            if not name:
                st.error("Nhap ten nganh.")
            elif name in sectors:
                st.error(f"Da co nganh \"{name}\" - sua o the \"Sua / xoa nganh\".")
            elif not cleaned:
                st.error("Nhap it nhat 1 ma co phieu.")
            else:
                st.error(validate_and_save({**sectors, name: cleaned}, cleaned,
                                           f"Da them nganh {name} ({len(cleaned)} ma)"))

    with tab_excel:
        st.markdown("**Xuat ra Excel**")
        st.download_button(
            "📥 Xuat Excel", data=sectors_to_excel(sectors),
            file_name=f"sectors_{datetime.now():%Y%m%d}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            on_click="ignore", key="mg_export",
        )
        st.caption(f"{len(sectors)} nganh, {sum(map(len, sectors.values()))} ma · Moi dong 1 ma, 2 cot: Sector | Ticker")

        st.markdown("**Nhap tu Excel**")
        uploaded = st.file_uploader("Chon file .xlsx (cot A: Sector, cot B: Ticker)", type=["xlsx"], key="mg_import_file")
        if uploaded is not None:
            try:
                imported = sectors_from_excel(uploaded)
            except ValueError as e:
                st.error(str(e))
            else:
                st.caption(f"Doc duoc {len(imported)} nganh, {sum(map(len, imported.values()))} ma: "
                           + ", ".join(imported))
                mode = st.segmented_control(
                    "Cach nhap", ["Thay the toan bo", "Gop vao danh sach hien tai"],
                    default="Thay the toan bo", key="mg_import_mode",
                )
                skip_unknown = st.checkbox("Bo qua ma khong tim thay tren bang gia", value=True, key="mg_import_skip")
                if st.button("📤 Nhap du lieu", type="primary", key="mg_import"):
                    if mode == "Gop vao danh sach hien tai":
                        result = {k: list(v) for k, v in sectors.items()}
                        for sector, tickers in imported.items():
                            bucket = result.setdefault(sector, [])
                            bucket.extend(t for t in tickers if t not in bucket)
                    else:
                        result = imported
                    known = {t for v in sectors.values() for t in v}
                    to_check = list(dict.fromkeys(t for v in result.values() for t in v if t not in known))
                    try:
                        unknown = find_unknown_tickers(to_check) if to_check else []
                    except requests.RequestException as e:
                        st.error(f"Khong kiem tra duoc ma co phieu (loi ket noi): {e}")
                    else:
                        if unknown and not skip_unknown:
                            st.error(f"Khong tim thay ma tren bang gia: {', '.join(unknown)}. "
                                     "Sua file, hoac tich \"Bo qua ma khong tim thay\".")
                        else:
                            result = {k: [t for t in v if t not in unknown] for k, v in result.items()}
                            result = {k: v for k, v in result.items() if v}  # bo nganh khong con ma nao
                            msg = f"Da nhap {len(result)} nganh, {sum(map(len, result.values()))} ma tu Excel"
                            if unknown:
                                msg += f" (bo qua: {', '.join(unknown)})"
                            st.error(validate_and_save(result, [], msg))

    st.divider()
    with st.popover("↩️ Khoi phuc danh sach mac dinh"):
        st.write("Bo moi thay doi va dung lai danh sach nganh ban dau?")
        if st.button("Xac nhan khoi phuc", key="mg_reset"):
            st.error(validate_and_save({k: list(v) for k, v in DEFAULT_SECTORS.items()}, [],
                                       "Da khoi phuc danh sach nganh mac dinh"))


# 7. Header Bar
col_h1, col_h2, col_h3, col_h4, col_h5, col_h6 = st.columns(
    [4.8, 1, 1, 1.15, 1.35, 1.2], vertical_alignment="center"
)
with col_h1:
    st.markdown('<div class="app-title">📈 Bang Gia Phan Nganh Chung Khoan</div>', unsafe_allow_html=True)
with col_h2:
    if st.button("⚙️ Quan ly nganh", width="stretch"):
        manage_sectors()
with col_h3:
    if st.button("🔄 Lam Moi Du Lieu", width="stretch"):
        store.refresh_now()
        get_simulated_data.clear()
with col_h4:
    with st.popover("📚 Stock history", width="stretch"):
        history_rows, latest_history_date = get_history_summary(STOCK_HISTORY_DATABASE)
        st.caption(
            f"{history_rows:,} daily rows"
            + (f" · latest: {latest_history_date}" if latest_history_date else "")
        )
        with st.form("stock_history_import"):
            history_files = st.file_uploader(
                "Import CafeF CSV files",
                type=["csv"],
                accept_multiple_files=True,
            )
            import_history = st.form_submit_button(
                "Import price history", width="stretch"
            )
        if import_history:
            if not history_files:
                st.error("Select at least one CSV file.")
            else:
                try:
                    import_progress = st.progress(
                        0.0, text="Preparing CSV import..."
                    )
                    imported_rows, skipped_rows = import_csv_files(
                        STOCK_HISTORY_DATABASE,
                        history_files,
                        progress_callback=lambda fraction, message: import_progress.progress(
                            min(fraction, 1.0), text=message
                        ),
                    )
                    history_rows, latest_history_date = get_history_summary(
                        STOCK_HISTORY_DATABASE
                    )
                    st.success(f"Imported or updated {imported_rows:,} rows.")
                    if skipped_rows:
                        st.warning(
                            f"Skipped {skipped_rows:,} rows without a ticker symbol."
                        )
                    st.caption(
                        f"Database total: {history_rows:,} rows · latest: "
                        f"{latest_history_date}"
                    )
                except Exception as error:
                    logging.exception("Could not import stock history CSV files")
                    st.error(f"CSV import failed: {error}")

        st.caption(
            "Daily close is captured after 15:00 on weekdays when this app is running "
            "and the market session is ended."
        )
        recent_signals = get_recent_signals(STOCK_HISTORY_DATABASE)
        if recent_signals:
            signal_rows = [
                {
                    "Ticker": item["ticker"],
                    "Date": item["trade_date"],
                    "Signal": item["signal"],
                    f"MA{item['ma_short_period']}": item["ma20"],
                    f"MA{item['ma_long_period']}": item["ma50"],
                    "RSI14": item["rsi14"],
                }
                for item in recent_signals
            ]
            st.dataframe(
                signal_rows,
                hide_index=True,
                alt="Recent BUY and SELL signals calculated from daily price history",
            )
        else:
            st.caption("No BUY or SELL crossover signals recorded yet.")
        (
            display_buy_sell_enabled,
            display_ma_short_period,
            display_ma_long_period,
            display_rsi_level1,
            display_rsi_level2,
        ) = store.signal_settings()
        st.caption(
            (
                f"BUY/SELL rule: MA{display_ma_short_period} crosses "
                f"MA{display_ma_long_period}, RSI "
                f"{display_rsi_level1:g}/{display_rsi_level2:g}."
            )
            if display_buy_sell_enabled
            else "BUY/SELL Telegram sending is disabled."
        )

with col_h5:
    with st.popover("📈 Tín hiệu", width="stretch"):
        st.subheader("Thiết lập tín hiệu MUA/BÁN")
        (
            current_buy_sell_enabled,
            current_ma_short,
            current_ma_long,
            current_rsi1,
            current_rsi2,
        ) = store.signal_settings()
        with st.form("buy_sell_signal_settings"):
            buy_sell_enabled_input = st.checkbox(
                "Gửi tín hiệu BUY/SELL qua Telegram",
                value=current_buy_sell_enabled,
            )
            ma_short_input = st.number_input(
                "MA ngắn hạn (số phiên)",
                min_value=2,
                max_value=499,
                value=int(current_ma_short),
                step=1,
            )
            ma_long_input = st.number_input(
                "MA dài hạn (số phiên)",
                min_value=3,
                max_value=500,
                value=int(current_ma_long),
                step=1,
            )
            rsi_level1_input = st.number_input(
                "RSI mức 1",
                min_value=0.0,
                max_value=99.0,
                value=float(current_rsi1),
                step=1.0,
            )
            rsi_level2_input = st.number_input(
                "RSI mức 2",
                min_value=1.0,
                max_value=100.0,
                value=float(current_rsi2),
                step=1.0,
            )
            st.caption(
                "BUY khi MA ngắn cắt lên MA dài và RSI nằm trong mức 1–2. "
                "SELL khi MA ngắn cắt xuống MA dài và RSI nằm trong "
                "(100 - mức 2) đến mức 1."
            )
            save_buy_sell_settings = st.form_submit_button(
                "Lưu thiết lập tín hiệu", width="stretch"
            )

        if save_buy_sell_settings:
            if not 2 <= ma_short_input < ma_long_input <= 500:
                st.error("MA ngắn hạn phải nhỏ hơn MA dài hạn (2 đến 500 phiên).")
            elif (
                not math.isfinite(rsi_level1_input)
                or not math.isfinite(rsi_level2_input)
                or not 0 <= rsi_level1_input < rsi_level2_input <= 100
            ):
                st.error("RSI mức 1 phải nhỏ hơn RSI mức 2 (trong khoảng 0–100).")
            else:
                (
                    current_alert_enabled,
                    current_threshold,
                    current_duration,
                    current_interval,
                    current_chat_ids,
                ) = store.alert_settings()
                current_price_update_enabled, _ = store.price_update_settings()
                try:
                    save_telegram_alert_settings(
                        current_alert_enabled,
                        current_threshold,
                        current_duration,
                        current_interval,
                        store.volume_settings(),
                        current_price_update_enabled,
                        buy_sell_enabled_input,
                        ma_short_input,
                        ma_long_input,
                        rsi_level1_input,
                        rsi_level2_input,
                        current_chat_ids,
                    )
                except OSError as error:
                    st.error(f"Could not save signal settings: {error}")
                else:
                    store.update_signal_settings(
                        buy_sell_enabled_input,
                        ma_short_input,
                        ma_long_input,
                        rsi_level1_input,
                        rsi_level2_input,
                    )
                    st.success("Đã lưu thiết lập tín hiệu BUY/SELL.")

with col_h6:
    with st.popover("📨 Telegram", width="stretch"):
        bot_status = get_telegram_bot_status()
        status_led = "🟢" if bot_status == "Running" else "🔴"
        st.caption(f"{status_led} Telegram command bot: {bot_status}")
        if st.button(
            "Stop Telegram Bot" if bot_status == "Running" else "Start Telegram Bot",
            width="stretch",
            key="telegram_bot_toggle",
        ):
            succeeded, message = (
                stop_telegram_bot()
                if bot_status == "Running"
                else start_telegram_bot()
            )
            if succeeded:
                st.success(message)
            else:
                st.error(message)
            st.caption("Startup output is written to telegram_bot.log.")

        with st.expander("User conversations", expanded=False):
            render_telegram_inbox()

        (
            current_enabled,
            current_threshold,
            current_duration,
            current_signal_interval,
            current_chat_ids,
        ) = store.alert_settings()
        current_price_update_enabled, _ = store.price_update_settings()
        current_volume_lookback = store.volume_settings()
        current_buy_sell_settings = store.signal_settings()
        with st.form("telegram_alert_settings"):
            enabled_input = st.checkbox(
                "Enable scheduled threshold snapshots",
                value=current_enabled,
            )
            threshold_input = st.number_input(
                "Daily change threshold (%)",
                min_value=0.1,
                max_value=100.0,
                value=float(current_threshold),
                step=0.1,
                format="%.2f",
            )
            signal_interval_input = st.number_input(
                "BUY/SELL signal interval (minutes)",
                min_value=5,
                max_value=360,
                value=int(current_signal_interval),
                step=5,
            )
            volume_lookback_input = st.selectbox(
                "Vol bình quân so với N phiên trước",
                options=list(range(10, 101, 5)),
                index=list(range(10, 101, 5)).index(current_volume_lookback),
                format_func=lambda sessions: f"{sessions} phiên",
            )
            st.caption(
                "Tỷ lệ = volume hiện tại / volume bình quân các phiên trước × 100%. "
                "Thiếu lịch sử sẽ hiển thị N/A."
            )
            st.caption(
                f"Threshold snapshots include all tickers at or beyond +/-{current_threshold:g}% "
                f"on weekdays at {PRICE_UPDATE_SCHEDULE_TEXT} (app local time)."
            )
            st.caption(
                f"Price changes are checked on weekdays at {PRICE_UPDATE_SCHEDULE_TEXT} "
                "(app local time). A summary is sent only when prices changed."
            )
            price_update_enabled_input = st.checkbox(
                "Send consolidated price changes to Telegram",
                value=current_price_update_enabled,
            )
            chat_ids_input = st.text_input(
                "Telegram chat IDs (comma-separated)",
                value=", ".join(current_chat_ids),
            )
            st.caption("For a private chat, send /chatid to the bot and enter the numeric ID it returns.")
            save_alert_settings = st.form_submit_button("Save alert settings", width="stretch")
            send_test = st.form_submit_button("Send Test", width="stretch")

        if save_alert_settings:
            chat_ids = parse_telegram_chat_ids(chat_ids_input)
            if not math.isfinite(threshold_input) or not 0 < threshold_input <= 100:
                st.error("Threshold must be between 0.1% and 100%.")
            elif not 5 <= signal_interval_input <= 360:
                st.error("Signal interval must be between 5 and 360 minutes.")
            else:
                try:
                    save_telegram_alert_settings(
                        enabled_input,
                        threshold_input,
                        current_duration,
                        signal_interval_input,
                        volume_lookback_input,
                        price_update_enabled_input,
                        *current_buy_sell_settings,
                        chat_ids,
                    )
                except OSError as e:
                    st.error(f"Could not save Telegram settings: {e}")
                else:
                    store.update_alert_settings(
                        enabled_input,
                        threshold_input,
                        current_duration,
                        signal_interval_input,
                        chat_ids,
                    )
                    store.update_price_update_settings(price_update_enabled_input)
                    store.update_volume_settings(volume_lookback_input)
                    current_enabled = enabled_input
                    current_threshold = threshold_input
                    current_signal_interval = signal_interval_input
                    current_volume_lookback = volume_lookback_input
                    current_price_update_enabled = price_update_enabled_input
                    current_chat_ids = chat_ids
                    st.success("Telegram alert settings saved.")

        if send_test:
            chat_ids = parse_telegram_chat_ids(chat_ids_input)
            if not TELEGRAM_BOT_TOKEN:
                st.error("Add TELEGRAM_BOT_TOKEN to Phong_telegram.env before sending a test.")
            elif not chat_ids:
                st.error("Enter at least one Telegram chat ID.")
            elif (
                not math.isfinite(threshold_input)
                or not 0 < threshold_input <= 100
                or not 5 <= signal_interval_input <= 360
                or volume_lookback_input not in range(10, 101, 5)
            ):
                st.error("Enter valid alert settings before sending a test.")
            else:
                quotes, quotes_updated_at, _, _ = store.snapshot()
                successful, failed = send_telegram_test(
                    chat_ids,
                    threshold_input,
                    current_price_update_enabled,
                    signal_interval_input,
                    volume_lookback_input,
                    current_buy_sell_settings,
                    list(store.tickers),
                    quotes,
                    quotes_updated_at,
                )
                if successful:
                    st.success(f"Test message sent to {len(successful)} recipient(s).")
                if failed:
                    failures = "; ".join(f"{chat_id}: {reason}" for chat_id, reason in failed)
                    st.error(f"Could not send test message: {failures}")

        if not TELEGRAM_BOT_TOKEN:
            st.warning("Add TELEGRAM_BOT_TOKEN to Phong_telegram.env to enable alerts.")
        elif not current_enabled:
            st.caption("Telegram price alerts are disabled.")
        elif not current_chat_ids:
            st.caption("No recipient IDs configured; alerts are paused.")

if "mg_flash" in st.session_state:
    st.toast(st.session_state.pop("mg_flash"), icon="✅")


# 8. Render Grid Layout - Fragment tu chay lai moi POLL_SECONDS giay, khong tai lai toan trang
@st.fragment(run_every=POLL_SECONDS)
def render_board():
    quotes, last_update, session, error = store.snapshot()
    now = datetime.now()
    board_frozen = is_board_frozen(now)
    updated = last_update or now

    if quotes:
        if board_frozen:
            source = "🔒 Giá chốt phiên / dữ liệu SQLite · dừng cập nhật đến 08:45 ngày làm việc tiếp theo"
        else:
            source = f"🟢 Realtime (Vietcap) · {SESSION_LABELS.get(session, session or 'N/A')}"
        if error and not board_frozen:
            source = f"🟡 Mat ket noi, dang thu lai - hien thi du lieu luc {updated:%H:%M:%S}"
    else:
        quotes = get_simulated_data()
        source = "🔴 Chua ket noi duoc nguon realtime - dang dung du lieu gia lap"
        if error:
            source += f" ({error[:80]})"
        updated = datetime.now()

    update_label = "Hien thi luc" if board_frozen else "Cap nhat luc"
    st.caption(f"{update_label}: {updated:%H:%M:%S} | Nguon: {source} | Tu lam moi moi {POLL_SECONDS}s"
               f" | Gia: nghin dong · KL: nghin co phieu | Telegram bot: {get_telegram_bot_status()}")

    if not SECTOR_MAP:
        st.info("Chua co nganh nao. Bam \"⚙️ Quan ly nganh\" de them.")
        return

    # Them cot khi danh sach dai ra (moi the nganh ~ so ma + 2 dong tieu de; ~32 dong/cot vua man hinh 1080p)
    row_units = sum(len(tickers) + 2 for tickers in SECTOR_MAP.values())
    column_count = min(8, max(6, math.ceil(row_units / 32)))

    # Toan bo cac nganh trong 1 khoi HTML, chia cot bang CSS (.board) de vua 1 man hinh
    # HTML phai viet lien, khong thut le/dong trong, neu khong Markdown se hieu la code block
    board_html = f'<div class="board" style="column-count: {column_count};">'
    for sector_name, tickers in SECTOR_MAP.items():
        board_html += (
            '<div class="sector-card">'
            '<div class="sector-header">'
            f'<span>{html.escape(sector_name)}</span>'
            '<span style="font-size: 10px; color: #8b949e;">MA | GIA | % | KL</span>'
            '</div>'
        )
        # Sap xep trong tung nganh theo % thay doi, tu cao xuong thap
        rows = sorted(((t, quotes[t]) for t in tickers if t in quotes), key=lambda r: r[1]["pct"], reverse=True)
        for ticker, s in rows:
            sign = "+" if s['pct'] > 0 else ""
            # Bam vao ma -> mo trang chi tiet tren 24hmoney o tab moi
            board_html += (
                f'<a class="stock-row {s["cls"]}" href="{STOCK_DETAIL_URL.format(ticker=ticker)}" '
                f'target="_blank" rel="noopener noreferrer" title="Xem {ticker} tren 24hmoney">'
                f'<span style="width: 25%;">{ticker}</span>'
                f'<span style="width: 25%;">{s["price"]:.2f}</span>'
                f'<span style="width: 30%;">{sign}{s["pct"]:.2f}%</span>'
                f'<span style="width: 20%; text-align: right; color: #8b949e;">{s["vol"]:,.0f}</span>'
                '</a>'
            )
        board_html += "</div>"
    board_html += "</div>"
    st.markdown(board_html, unsafe_allow_html=True)


render_board()
